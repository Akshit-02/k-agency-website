# Builds every calendar deliverable from posts_data.py and validates them.
# Usage (from this folder): python3 build.py   (needs openpyxl and network access for link checks)
import csv, datetime as dt, json, os, re, struct, urllib.request
import posts_data as D

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)                      # content/linkedin/30-day-launch
REPO = os.path.abspath(os.path.join(ROOT, "..", "..", ".."))
IMG_DIR = os.path.join(ROOT, "images")
BASE = "https://www.kudozz.in/blog/"
IST = dt.timezone(dt.timedelta(hours=5, minutes=30))
STATUS = "Not scheduled: ready for review"

# ------------------------------------------------------------------ assemble
def png_size(path):
    with open(path, "rb") as f:
        head = f.read(24)
    return struct.unpack(">II", head[16:24])

start = dt.date.fromisoformat(D.START_DATE)
register = {r["slug"]: r for r in csv.DictReader(open(os.path.join(REPO, "docs", "seo-register-blog-2026-10.csv"), encoding="utf-8"))}

posts = []
for i, p in enumerate(D.POSTS, 1):
    pid = f"KDZ-LI-{i:02d}"
    canonical = BASE + p["slug"]
    tracked = f"{canonical}?{D.UTM}&utm_content=post-{i:02d}"
    body = p["copy"].replace("{link}", tracked)
    final = body.rstrip() + "\n\n" + " ".join(p["hashtags"])
    date = start + dt.timedelta(days=i - 1)
    img = p.get("image", "")
    posts.append({
        "post_id": pid,
        "post_number": i,
        "date": date.isoformat(),
        "day": date.strftime("%A"),
        "time_ist": p["time"],
        "datetime_ist": f"{date.isoformat()}T{p['time']}:00+05:30",
        "category": p["category"],
        "objective": p["objective"],
        "target_audience": p["audience"],
        "hook": final.split("\n", 1)[0].strip(),
        "post_copy": final,
        "character_count": len(final),
        "blog_title": D.BLOG_TITLES[p["slug"]],
        "canonical_url": canonical,
        "tracked_url": tracked,
        "cta": p["cta"],
        "hashtags": p["hashtags"],
        "format": "Image" if img else "Text-only",
        "image_filename": img,
        "image_alt_text": p.get("alt", ""),
        "kudozz_service": D.SERVICES[p["service"]],
        "kudozz_service_url": "https://www.kudozz.in/services/" + p["service"],
        "scheduling_status": STATUS,
        "notes": p.get("notes", ""),
    })

# ------------------------------------------------------------------ link checks
class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *a, **k):
        return None
opener = urllib.request.build_opener(NoRedirect)

def check(url):
    req = urllib.request.Request(url, headers={"User-Agent": "kudozz-linkedin-calendar-check"})
    try:
        r = opener.open(req, timeout=30)
        page = r.read().decode("utf-8", "replace")
        m = re.search(r'<link rel="canonical" href="([^"]+)"', page)
        return r.status, (m.group(1) if m else "")
    except urllib.error.HTTPError as e:
        return e.code, ""
    except Exception as e:
        return f"error: {e}", ""

link_results = {}
for p in posts:
    s1, canon = check(p["canonical_url"])
    s2, canon2 = check(p["tracked_url"])
    ok = s1 == 200 and s2 == 200 and canon == p["canonical_url"] and canon2 == p["canonical_url"]
    link_results[p["canonical_url"]] = dict(status=s1, tracked_status=s2, canonical=canon, ok=ok)
    p["url_verification"] = "Verified" if ok else f"CHECK: status {s1}/{s2}, canonical {canon or 'missing'}"

# ------------------------------------------------------------------ validation
errors = []
now = dt.datetime.now(IST)
seen_slugs, seen_hooks = set(), set()
for p in posts:
    c, pid = p["post_copy"], p["post_id"]
    if p["tracked_url"] not in c: errors.append(f"{pid}: blog link missing from copy")
    if "{" in c or "[insert" in c.lower(): errors.append(f"{pid}: placeholder left in copy")
    if "—" in c or "–" in c: errors.append(f"{pid}: em/en dash in copy")
    if not 3 <= len(p["hashtags"]) <= 5: errors.append(f"{pid}: {len(p['hashtags'])} hashtags")
    if not c.rstrip().endswith(" ".join(p["hashtags"])): errors.append(f"{pid}: hashtags in copy don't match metadata")
    if len(set(h.lower() for h in p["hashtags"])) != len(p["hashtags"]): errors.append(f"{pid}: duplicate hashtag")
    for h in p["hashtags"]:
        if not re.fullmatch(r"#[A-Za-z][A-Za-z0-9]+", h): errors.append(f"{pid}: malformed hashtag {h}")
    if len(c) > 3000: errors.append(f"{pid}: {len(c)} chars exceeds LinkedIn's 3,000 limit")
    if p["canonical_url"] in seen_slugs: errors.append(f"{pid}: blog reused")
    seen_slugs.add(p["canonical_url"])
    if p["hook"] in seen_hooks: errors.append(f"{pid}: duplicate hook")
    seen_hooks.add(p["hook"])
    if dt.datetime.fromisoformat(p["datetime_ist"]) <= now: errors.append(f"{pid}: date is in the past")
    if not p["url_verification"] == "Verified": errors.append(f"{pid}: {p['url_verification']}")
    if p["image_filename"]:
        path = os.path.join(IMG_DIR, p["image_filename"])
        if not os.path.exists(path): errors.append(f"{pid}: image file missing")
        elif png_size(path) != (1200, 1200): errors.append(f"{pid}: image not 1200x1200")
        if not p["image_filename"].startswith(f"post-{p['post_number']:02d}-"): errors.append(f"{pid}: image name doesn't match post")
        if len(p["image_alt_text"]) < 40: errors.append(f"{pid}: alt text missing or too short")
img_posts = [p for p in posts if p["image_filename"]]
on_disk = sorted(f for f in os.listdir(IMG_DIR) if f.endswith(".png"))
if sorted(p["image_filename"] for p in img_posts) != on_disk: errors.append("image folder doesn't match the calendar")
if len(posts) != 30: errors.append(f"{len(posts)} posts, expected 30")
if not 10 <= len(img_posts) <= 12: errors.append(f"{len(img_posts)} image posts")
dates = [p["date"] for p in posts]
if len(set(dates)) != 30: errors.append("dates are not 30 consecutive days")

# ------------------------------------------------------------------ writers
COLS = [("post_id", "Post ID"), ("post_number", "Post number"), ("date", "Date"), ("day", "Day"), ("time_ist", "Time (IST)"),
        ("category", "Category"), ("objective", "Objective"), ("target_audience", "Target audience"), ("hook", "Hook"),
        ("post_copy", "Full post copy"), ("character_count", "Characters"), ("blog_title", "Blog title"),
        ("canonical_url", "Canonical blog URL"), ("tracked_url", "Tracked URL"), ("url_verification", "URL verification"),
        ("cta", "CTA"), ("hashtags", "Hashtags"), ("format", "Image or text-only"), ("image_filename", "Image filename"),
        ("image_alt_text", "Alt text"), ("kudozz_service", "Kudozz service"), ("scheduling_status", "Manual scheduling status"),
        ("notes", "Notes")]

def cell(p, k):
    v = p[k]
    return " ".join(v) if isinstance(v, list) else v

def write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, quoting=csv.QUOTE_ALL, lineterminator="\r\n")
        w.writerow(header)
        w.writerows(rows)

write_csv(os.path.join(ROOT, "linkedin-30-day-calendar.csv"), [h for _, h in COLS], [[cell(p, k) for k, _ in COLS] for p in posts])

with open(os.path.join(ROOT, "linkedin-posts.json"), "w", encoding="utf-8") as f:
    json.dump({"campaign": "Kudozz LinkedIn 30-day launch", "timezone": "Asia/Kolkata (IST, UTC+05:30)",
               "date_range": [posts[0]["date"], posts[-1]["date"]], "generated": now.isoformat(timespec="minutes"),
               "utm": D.UTM + "&utm_content=post-NN", "posts": posts}, f, ensure_ascii=False, indent=2)

write_csv(os.path.join(ROOT, "image-manifest.csv"),
          ["Post ID", "Date", "Image filename", "Post topic", "Alt text", "Image format", "Image dimensions", "Asset status"],
          [[p["post_id"], p["date"], p["image_filename"], p["blog_title"], p["image_alt_text"], "PNG",
            "%dx%d" % png_size(os.path.join(IMG_DIR, p["image_filename"])), "Created; ready for review"] for p in img_posts])

write_csv(os.path.join(ROOT, "blog-references.csv"),
          ["Blog title", "Canonical URL", "Associated post IDs", "Published", "Last updated", "Main topic (site category)",
           "Verification status", "Checked (IST)", "Notes"],
          [[p["blog_title"], p["canonical_url"], p["post_id"], register.get(p["canonical_url"][len(BASE):], {}).get("published", ""),
            register.get(p["canonical_url"][len(BASE):], {}).get("updated", ""), register.get(p["canonical_url"][len(BASE):], {}).get("category", ""),
            p["url_verification"], now.strftime("%Y-%m-%d %H:%M"),
            "HTTP 200 with no redirect; self-referencing canonical; tracked URL also returns 200 and canonicalises to the clean URL."
            if p["url_verification"] == "Verified" else "Needs review before posting"] for p in posts])

# Markdown
md = [f"# Kudozz LinkedIn calendar: 30-day launch",
      "",
      f"{posts[0]['date']} to {posts[-1]['date']}. All times are IST (Asia/Kolkata). Nothing here has been scheduled or published.",
      "",
      "Each post's copy block below is the exact text to paste into LinkedIn, including the blog link and hashtags.",
      "",
      "## At a glance",
      "",
      "| # | Date | Time | Category | Format | Blog |",
      "|---|---|---|---|---|---|"]
for p in posts:
    md.append(f"| [{p['post_number']:02d}](#{p['post_id'].lower()}) | {p['day'][:3]} {p['date']} | {p['time_ist']} | {p['category']} | {p['format']} | {p['blog_title']} |")
for p in posts:
    md += ["", "---", "", f"## {p['post_id']}",
           "", f"**{p['day']} {p['date']}, {p['time_ist']} IST** · {p['category']} · {p['format']}", "",
           f"- **Objective:** {p['objective']}",
           f"- **Audience:** {p['target_audience']}",
           f"- **Blog:** [{p['blog_title']}]({p['canonical_url']}) ({p['url_verification']})",
           f"- **CTA:** {p['cta']}",
           f"- **Hashtags:** {' '.join(p['hashtags'])}",
           f"- **Kudozz service:** {p['kudozz_service']}",
           f"- **Status:** {p['scheduling_status']}"]
    if p["image_filename"]:
        md += [f"- **Image:** `images/{p['image_filename']}`", f"- **Alt text:** {p['image_alt_text']}"]
    if p["notes"]:
        md.append(f"- **Notes:** {p['notes']}")
    md += ["", f"**Copy ({p['character_count']} characters):**", "", "```text", p["post_copy"], "```"]
with open(os.path.join(ROOT, "linkedin-30-day-calendar.md"), "w", encoding="utf-8") as f:
    f.write("\n".join(md) + "\n")

# XLSX
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
WIDTH = {"post_copy": 80, "image_alt_text": 50, "objective": 40, "target_audience": 34, "hook": 50, "blog_title": 40, "canonical_url": 46,
         "tracked_url": 50, "cta": 38, "hashtags": 34, "notes": 36, "category": 26, "kudozz_service": 28, "scheduling_status": 26,
         "image_filename": 38, "url_verification": 14}
wb = Workbook(); ws = wb.active; ws.title = "Calendar"
ws.append([h for _, h in COLS])
for p in posts:
    row = []
    for k, _ in COLS:
        v = cell(p, k)
        if k == "date": v = dt.date.fromisoformat(v)
        if k == "time_ist": v = dt.time.fromisoformat(v)
        row.append(v)
    ws.append(row)
head_fill = PatternFill("solid", fgColor="150F0B"); thin = Side(style="thin", color="E8D8B8")
for ci, (k, _) in enumerate(COLS, 1):
    letter = ws.cell(1, ci).column_letter
    ws.column_dimensions[letter].width = WIDTH.get(k, 13)
    h = ws.cell(1, ci); h.font = Font(bold=True, color="FAF1E2"); h.fill = head_fill
    h.alignment = Alignment(vertical="center", wrap_text=True)
    for ri in range(2, len(posts) + 2):
        c = ws.cell(ri, ci)
        c.alignment = Alignment(vertical="top", wrap_text=k not in ("canonical_url", "tracked_url", "image_filename"))
        c.border = Border(bottom=thin)
        if k == "date": c.number_format = "ddd dd mmm yyyy"
        if k == "time_ist": c.number_format = "hh:mm"
        if k in ("canonical_url", "tracked_url"): c.hyperlink = c.value; c.font = Font(color="5A2EEA", underline="single")
for ri in range(2, len(posts) + 2):
    ws.row_dimensions[ri].height = 409  # Excel's max; post copy cells scroll within
ws.row_dimensions[1].height = 32
ws.freeze_panes = "C2"
ws.auto_filter.ref = ws.dimensions
status_col = [i for i, (k, _) in enumerate(COLS, 1) if k == "scheduling_status"][0]
dv = DataValidation(type="list", formula1='"Not scheduled: ready for review,Approved,Scheduled in LinkedIn,Published,On hold"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"{ws.cell(2, status_col).column_letter}2:{ws.cell(31, status_col).column_letter}31")
info = wb.create_sheet("Read me")
for line in ["Kudozz LinkedIn 30-day launch calendar", f"Dates: {posts[0]['date']} to {posts[-1]['date']} (IST, Asia/Kolkata)",
             "Copy the 'Full post copy' cell as-is: it already contains the tracked blog link and hashtags.",
             "Update 'Manual scheduling status' from the dropdown as you schedule and publish.",
             "Images are in the images/ folder next to this workbook. Add the alt text when uploading.",
             "Suggested times are recommendations only; nothing has been scheduled."]:
    info.append([line])
info.column_dimensions["A"].width = 110; info["A1"].font = Font(bold=True, size=14)
xlsx = os.path.join(ROOT, "linkedin-30-day-calendar.xlsx")
wb.save(xlsx)

# ------------------------------------------------------------------ cross-file checks
wb2 = load_workbook(xlsx); ws2 = wb2["Calendar"]
xl_rows = list(ws2.iter_rows(min_row=2, values_only=True))
copy_idx = [k for k, _ in COLS].index("post_copy")
if len(xl_rows) != 30 or [r[copy_idx] for r in xl_rows] != [p["post_copy"] for p in posts]: errors.append("XLSX copy differs from source")
with open(os.path.join(ROOT, "linkedin-30-day-calendar.csv"), encoding="utf-8-sig", newline="") as f:
    rows = list(csv.reader(f))[1:]
if len(rows) != 30 or [r[copy_idx] for r in rows] != [p["post_copy"] for p in posts]: errors.append("CSV copy differs from source")
j = json.load(open(os.path.join(ROOT, "linkedin-posts.json"), encoding="utf-8"))
if [p["post_copy"] for p in j["posts"]] != [p["post_copy"] for p in posts]: errors.append("JSON copy differs from source")
mdtext = open(os.path.join(ROOT, "linkedin-30-day-calendar.md"), encoding="utf-8").read()
if not all(p["post_copy"] in mdtext for p in posts): errors.append("Markdown missing a post's full copy")

# ------------------------------------------------------------------ report
from collections import Counter
print("categories:", dict(Counter(p["category"] for p in posts)))
print("image posts:", len(img_posts), "text-only:", 30 - len(img_posts))
print("unique blog URLs:", len(seen_slugs), "verified:", sum(p["url_verification"] == "Verified" for p in posts))
print("hashtags per post:", dict(Counter(len(p["hashtags"]) for p in posts)))
print("top hashtags:", Counter(h for p in posts for h in p["hashtags"]).most_common(6))
print("chars min/max:", min(p["character_count"] for p in posts), max(p["character_count"] for p in posts))
print("range:", posts[0]["datetime_ist"], "->", posts[-1]["datetime_ist"])
print("ERRORS:" if errors else "ALL CHECKS PASSED", *errors, sep="\n  ")
